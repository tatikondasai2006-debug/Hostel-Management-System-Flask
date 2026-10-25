"""Generate a clean, professional deck for the hostel app.

Inspired by "Create a Professional PowerPoint Welcome Slide Step-by-Step":
a light, minimal, high-contrast look (white ground, one confident accent
colour, generous spacing, a clear visual hierarchy) opened by a polished
Welcome slide, then the same pin-to-pin workflows presented step by step.

Output: ``Hostel_Professional_Deck.pptx`` in the project root.
Every slide carries a Morph transition (with a fade fallback).
"""

from __future__ import annotations

import contextlib
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUT = PROJECT_ROOT / "Hostel_Professional_Deck.pptx"

# Your photo, circular-cropped into the welcome slide's round space.
IMAGE_SRC = Path(r"C:\Users\leela\OneDrive\Videos\Pictures\mywallpaper.jpeg")
CIRCLE_PNG = PROJECT_ROOT / "deck_assets" / "welcome_circle.png"

# Auto-advance: the show moves on by itself (no mouse click), Morph between
# slides. Seconds to linger on each slide type (roughly how long it reads).
MS_WELCOME = 6000
MS_SECTION = 6000
MS_MAP = 5000
MS_STEP = 8000
MS_LIFE = 6000
MS_GUAR = 7000

# --- palette (light, professional) -----------------------------------------
WHITE = "FFFFFF"
SOFT = "F1F5FB"  # page tint / muted panels
SOFT2 = "E8EEF7"  # slightly deeper panel
INK = "0F172A"  # primary text
SUB = "475569"  # secondary text
MUTED = "94A3B8"  # tertiary text
LINE = "DCE4F0"  # hairline borders
ACC = "2563EB"  # primary accent (blue)
ACC2 = "0EA5E9"  # secondary accent (sky)
TEAL = "0D9488"
GREEN = "16A34A"
AMBER = "D97706"
VIOLET = "7C3AED"
FONT = "Segoe UI Semibold"
FONT_BODY = "Segoe UI"

HOSTEL_NAME = "Sunrise Student Hostel"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

ROLE_COLOR = {"ADMIN": ACC, "STUDENT": VIOLET, "ADMIN + STUDENT": AMBER}


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


# --- transitions -----------------------------------------------------------
def morph(slide, advance_ms=MS_SECTION):
    sld = slide._element
    for t in sld.findall(qn("p:transition")):
        sld.remove(t)
    xml = (
        "<p:transition "
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" '
        f'spd="slow" advClick="0" advTm="{advance_ms}"><mc:AlternateContent>'
        '<mc:Choice Requires="p159"><p159:morph option="byObject"/></mc:Choice>'
        "<mc:Fallback><p:fade/></mc:Fallback>"
        "</mc:AlternateContent></p:transition>"
    )
    ref = sld.find(qn("p:clrMapOvr"))
    if ref is None:
        ref = sld.find(qn("p:cSld"))
    ref.addnext(parse_xml(xml))


def make_circle_png(src=IMAGE_SRC, dst=CIRCLE_PNG, size=1000):
    """Centre-crop ``src`` to a square and mask it to a soft-edged circle."""
    if not Path(src).exists():
        return False
    from PIL import Image, ImageDraw

    dst.parent.mkdir(parents=True, exist_ok=True)
    im = Image.open(src).convert("RGBA")
    w, h = im.size
    s = min(w, h)
    box = ((w - s) // 2, (h - s) // 2, (w - s) // 2 + s, (h - s) // 2 + s)
    im = im.crop(box).resize((size, size), Image.LANCZOS)
    ss = 4  # supersample the mask for smooth edges
    mask = Image.new("L", (size * ss, size * ss), 0)
    ImageDraw.Draw(mask).ellipse((0, 0, size * ss - 1, size * ss - 1), fill=255)
    im.putalpha(mask.resize((size, size), Image.LANCZOS))
    im.save(dst)
    return True


# --- primitives ------------------------------------------------------------
def no_line(shape):
    shape.line.fill.background()


def shape(slide, kind, x, y, w, h, fill=None, line=None, lw=1.0, radius=0.14):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        no_line(shp)
    else:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        with contextlib.suppress(IndexError, ValueError):
            shp.adjustments[0] = radius
    shp.shadow.inherit = False
    return shp


def grad(slide, kind, x, y, w, h, c1, c2, angle=0.0, radius=0.2):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.gradient()
    with contextlib.suppress(Exception):
        shp.fill.gradient_angle = angle
    stops = shp.fill.gradient_stops
    stops[0].color.rgb = rgb(c1)
    stops[0].position = 0.0
    stops[-1].color.rgb = rgb(c2)
    stops[-1].position = 1.0
    no_line(shp)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        with contextlib.suppress(IndexError, ValueError):
            shp.adjustments[0] = radius
    shp.shadow.inherit = False
    return shp


def _fill(tf, lines):
    tf.clear()
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get("align", PP_ALIGN.LEFT)
        p.level = ln.get("level", 0)
        if "space_after" in ln:
            p.space_after = Pt(ln["space_after"])
        if "space_before" in ln:
            p.space_before = Pt(ln["space_before"])
        if "line_spacing" in ln:
            p.line_spacing = ln["line_spacing"]
        runs = ln.get("runs") or [{"text": ln.get("text", "")}]
        for r in runs:
            run = p.add_run()
            run.text = r.get("text", "")
            f = run.font
            f.size = Pt(r.get("size", ln.get("size", 14)))
            f.bold = r.get("bold", ln.get("bold", False))
            f.italic = r.get("italic", ln.get("italic", False))
            f.name = r.get("font", ln.get("font", FONT_BODY))
            f.color.rgb = rgb(r.get("color", ln.get("color", SUB)))


def txt(slide, x, y, w, h, lines, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    _fill(tf, lines)
    return tb


def stext(shp, lines, anchor=MSO_ANCHOR.MIDDLE, m=0.08):
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(m)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    _fill(tf, lines)


def bullets(items, size=13, color=SUB, gap=8, dot=ACC, font=FONT_BODY):
    return [
        {
            "runs": [
                {"text": "\u25aa  ", "color": dot, "size": size, "bold": True},
                {"text": t, "color": color, "size": size, "font": font},
            ],
            "space_after": gap,
        }
        for t in items
    ]


# --- slide chrome ----------------------------------------------------------
def photo_badge(slide, x, y, d):
    """Small circular photo accent (reuses the welcome-circle image)."""
    if not CIRCLE_PNG.exists():
        return
    shape(slide, MSO_SHAPE.OVAL, x - 0.03, y - 0.03, d + 0.06, d + 0.06, fill=WHITE)
    slide.shapes.add_picture(str(CIRCLE_PNG), Inches(x), Inches(y), Inches(d), Inches(d))
    shape(slide, MSO_SHAPE.OVAL, x - 0.03, y - 0.03, d + 0.06, d + 0.06, line=ACC, lw=1.0)


def new_slide(advance_ms=MS_SECTION, badge=True):
    s = prs.slides.add_slide(BLANK)
    morph(s, advance_ms)
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, fill=WHITE)
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.14, fill=ACC)
    shape(s, MSO_SHAPE.OVAL, 11.4, -1.2, 4.2, 4.2, fill=SOFT)
    shape(s, MSO_SHAPE.OVAL, -1.4, 5.6, 3.6, 3.6, fill=SOFT)
    if badge:
        photo_badge(s, 12.52, 0.24, 0.7)
    return s


def header(slide, kicker, title, sub=None):
    shape(slide, MSO_SHAPE.RECTANGLE, 0.7, 0.66, 0.1, 0.46, fill=ACC)
    txt(
        slide,
        0.7,
        0.6,
        11.5,
        0.28,
        [{"runs": [{"text": kicker.upper(), "size": 10.5, "bold": True, "color": ACC}]}],
    )
    txt(
        slide,
        0.7,
        0.88,
        11.9,
        0.6,
        [{"runs": [{"text": title, "size": 27, "bold": True, "color": INK, "font": FONT}]}],
    )
    if sub:
        txt(slide, 0.72, 1.55, 11.9, 0.4, [{"runs": [{"text": sub, "size": 13, "color": SUB}]}])
    shape(slide, MSO_SHAPE.RECTANGLE, 0.7, 2.02, 11.93, 0.02, fill=LINE)


def card(slide, x, y, w, h, radius=0.06, fill=WHITE):
    shape(
        slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=LINE, lw=1.0, radius=radius
    )


def chip(slide, x, y, w, h, label, color=ACC):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=color, radius=0.5)
    txt(
        slide,
        x,
        y,
        w,
        h,
        [
            {
                "align": PP_ALIGN.CENTER,
                "runs": [{"text": label, "size": 10, "bold": True, "color": WHITE}],
            }
        ],
        anchor=MSO_ANCHOR.MIDDLE,
    )


def rail(slide, nodes, current):
    txt(
        slide,
        0.7,
        0.32,
        8,
        0.24,
        [{"runs": [{"text": "STEP-BY-STEP WORKFLOW", "size": 9, "bold": True, "color": ACC}]}],
    )
    n = len(nodes)
    left, right, gap = 0.7, 12.63, 0.12
    w = (right - left - gap * (n - 1)) / n
    y, h = 0.64, 0.42
    for i, name in enumerate(nodes):
        x = left + i * (w + gap)
        if i == current:
            shp = shape(
                slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y - 0.05, w, h + 0.1, fill=ACC, radius=0.4
            )
            tc = WHITE
        else:
            done = i < current
            shp = shape(
                slide,
                MSO_SHAPE.ROUNDED_RECTANGLE,
                x,
                y,
                w,
                h,
                fill=(SOFT2 if done else WHITE),
                line=(ACC if done else LINE),
                lw=1.0,
                radius=0.4,
            )
            tc = ACC if done else MUTED
        stext(
            shp,
            [
                {
                    "align": PP_ALIGN.CENTER,
                    "runs": [
                        {"text": f"{i + 1} ", "size": 9, "bold": True, "color": tc},
                        {"text": name, "size": 9, "bold": (i == current), "color": tc},
                    ],
                }
            ],
        )


def info(slide, x, y, w, h, label, value, color=ACC):
    card(slide, x, y, w, h)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y + 0.12, 0.07, h - 0.24, fill=color)
    txt(
        slide,
        x + 0.26,
        y + 0.13,
        w - 0.42,
        h - 0.2,
        [
            {
                "runs": [{"text": label.upper(), "size": 9, "bold": True, "color": color}],
                "space_after": 3,
            },
            {"runs": [{"text": value, "size": 11.5, "color": INK}]},
        ],
    )


# --- content ---------------------------------------------------------------
ADMIN_NODES = [
    "Sign in",
    "Dashboard",
    "Hostels & Rooms",
    "Students",
    "Allocations",
    "Fees",
    "Payments",
    "Reports & Audit",
]
STUDENT_NODES = [
    "Sign up",
    "Sign in",
    "Dashboard",
    "Profile",
    "Allocations",
    "Fees",
    "Payments",
    "Receipts",
]

ADMIN_STEPS = [
    {
        "title": "Admin signs in",
        "route": "/auth/login \u2192 /admin/",
        "bullets": [
            "Email + password verified against the stored hash.",
            "A Flask-Login session opens and last_login_at is stamped.",
            "Role decides the landing page: admin \u2192 dashboard.",
        ],
        "files": "app/auth/routes.py",
        "data": "credentials \u2192 session",
        "guard": "rate-limited \u00b7 no enumeration \u00b7 role redirect",
    },
    {
        "title": "Operational dashboard",
        "route": "/admin/  (admin.dashboard)",
        "bullets": [
            "Live KPIs: students, rooms, capacity, occupancy.",
            "Monthly collection, outstanding and overdue count.",
            "Charts: by room type, by hostel, 6-month trend.",
        ],
        "files": "admin/routes.py \u00b7 services/dashboard_service.py",
        "data": "SQL aggregates \u2192 cards + charts",
        "guard": "admin_required",
    },
    {
        "title": "Set up hostels & rooms",
        "route": "/hostels \u00b7 /rooms",
        "bullets": [
            "Create a hostel, then add rooms (type, capacity, base fee).",
            "Capacity can never drop below current occupancy.",
            "Rooms with history can't be deleted \u2014 make unavailable.",
        ],
        "files": "rooms/routes.py \u00b7 rooms/forms.py",
        "data": "form \u2192 Hostel / Room rows",
        "guard": "capacity \u2265 occupancy \u00b7 unique number",
    },
    {
        "title": "Onboard students",
        "route": "/students",
        "bullets": [
            "Create a student + optional portal login in one step.",
            "Filter by course, year, status, allocation, fee state.",
            "Deactivating automatically releases the active room.",
        ],
        "files": "students/routes.py",
        "data": "profile \u2192 Student (+ User)",
        "guard": "unique email & roll no \u00b7 deactivation frees bed",
    },
    {
        "title": "Allocate beds",
        "route": "/allocations",
        "bullets": [
            "Auto, manual, transfer and release operations.",
            "Auto picks the lowest free bed in the first eligible room.",
            "Concurrency-safe: SKIP LOCKED + unique indexes + trigger.",
        ],
        "files": "allocations/routes.py \u00b7 services/allocation_service.py",
        "data": "student + room \u2192 active allocation",
        "guard": "no double-booking \u00b7 retries on contention",
    },
    {
        "title": "Generate monthly bills",
        "route": "/fees  (+ /fees/structures)",
        "bullets": [
            "Structures per room type set the default amounts.",
            "Generate for a period (YYYY-MM); duplicates skipped.",
            "Late fee after grace; status paid / partial / overdue / pending.",
        ],
        "files": "fees/routes.py \u00b7 services/fee_service.py",
        "data": "rooms + structures + period \u2192 StudentFee rows",
        "guard": "Decimal money \u00b7 idempotent generation",
    },
    {
        "title": "Record payments & receipts",
        "route": "/payments",
        "bullets": [
            "Idempotency key prevents duplicate payments.",
            "Fee row locked (FOR UPDATE); overpayment rejected.",
            "Immutable \u2014 corrections are linked reversals; PDF receipt.",
        ],
        "files": "payments/routes.py \u00b7 services/payment_service.py",
        "data": "payment \u2192 Payment row \u2192 balance update",
        "guard": "idempotent \u00b7 immutable \u00b7 receipt PDF",
    },
    {
        "title": "Reports, audit & settings",
        "route": "/reports \u00b7 /admin/audit \u00b7 /admin/settings",
        "bullets": [
            "Slug-based reports with CSV export + occupancy/financial PDFs.",
            "Audit log records who changed what.",
            "Settings: hostel identity, currency, late-fee policy.",
        ],
        "files": "reports/routes.py \u00b7 admin/routes.py",
        "data": "DB \u2192 CSV / PDF \u00b7 every write \u2192 audit",
        "guard": "admin_required \u00b7 immutable audit trail",
    },
]

STUDENT_STEPS = [
    {
        "title": "Create a portal account",
        "route": "/auth/signup",
        "bullets": [
            "Signup creates the User (role = student) + Student profile.",
            "Email and roll number must be unique.",
            "New accounts link straight to the student portal.",
        ],
        "files": "app/auth/routes.py",
        "data": "signup \u2192 User + Student",
        "guard": "rate-limited \u00b7 unique email & roll no",
    },
    {
        "title": "Student signs in",
        "route": "/auth/login \u2192 /portal",
        "bullets": [
            "Same login endpoint; the role picks the destination.",
            "Students always land on the portal dashboard.",
        ],
        "files": "app/auth/routes.py",
        "data": "credentials \u2192 session",
        "guard": "rate-limited \u00b7 role redirect",
    },
    {
        "title": "Personal dashboard",
        "route": "/portal",
        "bullets": [
            "Current room allocation and recent payments.",
            "Upcoming due bill and totals: billed, paid, outstanding.",
            "Everything scoped to the logged-in student.",
        ],
        "files": "students/portal.py",
        "data": "own records \u2192 summary",
        "guard": "student_required \u00b7 self-scoped",
    },
    {
        "title": "Profile self-service",
        "route": "/portal/profile",
        "bullets": [
            "View the profile; edit phone, address, emergency contact.",
            "Identity fields stay admin-controlled.",
        ],
        "files": "students/portal.py",
        "data": "limited form \u2192 Student",
        "guard": "editable fields restricted",
    },
    {
        "title": "My room",
        "route": "/portal/allocations",
        "bullets": [
            "Current bed plus full allocation history.",
            "Read-only \u2014 allocation is an admin action.",
        ],
        "files": "students/portal.py",
        "data": "own allocations \u2192 list",
        "guard": "self-scoped",
    },
    {
        "title": "My bills",
        "route": "/portal/fees",
        "bullets": [
            "Monthly bills with amount, paid, balance, status.",
            "Per-bill detail shows the breakdown and due date.",
            "Opening someone else's bill returns 404.",
        ],
        "files": "students/portal.py",
        "data": "own fees \u2192 list / detail",
        "guard": "IDOR-safe (404 on others)",
    },
    {
        "title": "My payments",
        "route": "/portal/payments",
        "bullets": [
            "History of payments recorded against the account.",
            "Read-only \u2014 recording is an admin action.",
        ],
        "files": "students/portal.py",
        "data": "own payments \u2192 list",
        "guard": "self-scoped",
    },
    {
        "title": "Download receipts",
        "route": "/portal/receipts/<id>",
        "bullets": [
            "Download a PDF receipt for any of your own payments.",
            "A receipt that isn't yours returns 404.",
        ],
        "files": "students/portal.py",
        "data": "payment id \u2192 own PDF",
        "guard": "IDOR-safe (404 on others)",
    },
]

STACK = [
    ("\u2699", "Flask 3", "App factory + blueprints"),
    ("\U0001f5c4", "PostgreSQL", "SQLAlchemy + Alembic"),
    ("\U0001f510", "Flask-Login", "Role-based sessions"),
    ("\U0001f310", "Bootstrap 5", "Responsive UI"),
    ("\U0001f4c4", "ReportLab", "PDF receipts & reports"),
    ("\u2713", "pytest", "147 tests passing"),
]

GUARANTEES = [
    ("\U0001f510", "Role-checked routes", "admin_required / student_required on every view"),
    ("\U0001f5c3", "DB-level integrity", "partial unique indexes + capacity triggers"),
    ("\U0001f4b3", "Immutable payments", "corrections are linked reversal records"),
    ("\U0001f464", "Self-scoped data", "students never see other records (404)"),
    ("\U0001f522", "Decimal money", "no floating-point drift in billing"),
    ("\U0001f4cb", "Full audit trail", "every state change is logged"),
]

LIFECYCLE = [
    ("Set up hostels\n& rooms", "ADMIN"),
    ("Enroll\nstudents", "ADMIN + STUDENT"),
    ("Allocate\nbeds", "ADMIN"),
    ("Generate\nmonthly fees", "ADMIN"),
    ("Record\npayments", "ADMIN"),
    ("View bills &\nreceipts", "STUDENT"),
    ("Reports, audit\n& deactivate", "ADMIN"),
]

FLOW = ["Setup", "Rooms", "Students", "Allocate", "Bill", "Pay", "Report"]


# --- slides ----------------------------------------------------------------
def s_welcome():
    s = new_slide(MS_WELCOME, badge=False)
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.18, fill=ACC)
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0.18, 4.2, 0.06, fill=ACC2)

    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.0, 1.35, 0.62, 0.62, fill=ACC, radius=0.25)
    txt(
        s,
        1.0,
        1.35,
        0.62,
        0.62,
        [
            {
                "align": PP_ALIGN.CENTER,
                "runs": [{"text": "H", "size": 26, "bold": True, "color": WHITE, "font": FONT}],
            }
        ],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    txt(
        s,
        1.78,
        1.42,
        8.5,
        0.5,
        [
            {
                "runs": [
                    {"text": "HOSTEL MANAGEMENT SYSTEM", "size": 12.5, "bold": True, "color": SUB}
                ]
            }
        ],
    )

    txt(
        s,
        1.0,
        2.35,
        8.6,
        0.4,
        [{"runs": [{"text": "WELCOME", "size": 13, "bold": True, "color": ACC}]}],
    )
    txt(
        s,
        1.0,
        2.85,
        8.6,
        1.7,
        [
            {
                "runs": [
                    {
                        "text": "Sunrise Student",
                        "size": 52,
                        "bold": True,
                        "color": INK,
                        "font": FONT,
                    }
                ],
                "line_spacing": 0.95,
                "space_after": 0,
            },
            {
                "runs": [{"text": "Hostel", "size": 52, "bold": True, "color": ACC, "font": FONT}],
                "line_spacing": 0.95,
            },
        ],
    )
    shape(s, MSO_SHAPE.RECTANGLE, 1.02, 4.85, 2.4, 0.06, fill=ACC2)
    txt(
        s,
        1.02,
        5.05,
        8.4,
        0.5,
        [
            {
                "runs": [
                    {
                        "text": "Admin & Student Portal \u2014 a clear, end-to-end "
                        "walkthrough of the whole operation.",
                        "size": 15,
                        "color": SUB,
                    }
                ]
            }
        ],
    )

    x = 1.02
    for label in FLOW:
        w = 0.32 + 0.13 * len(label)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 5.85, w, 0.4, fill=SOFT, radius=0.5)
        txt(
            s,
            x,
            5.85,
            w,
            0.4,
            [
                {
                    "align": PP_ALIGN.CENTER,
                    "runs": [{"text": label, "size": 10.5, "bold": True, "color": ACC}],
                }
            ],
            anchor=MSO_ANCHOR.MIDDLE,
        )
        x += w + 0.12

    txt(
        s,
        1.02,
        6.68,
        8.4,
        0.35,
        [
            {
                "runs": [
                    {
                        "text": "Flask  \u00b7  PostgreSQL  \u00b7  2026      Auto-advancing",
                        "size": 11,
                        "color": MUTED,
                    }
                ]
            }
        ],
    )
    shape(s, MSO_SHAPE.OVAL, 10.05, 1.45, 3.7, 3.7, fill=SOFT)
    if CIRCLE_PNG.exists():
        s.shapes.add_picture(str(CIRCLE_PNG), Inches(10.15), Inches(1.55), Inches(3.5), Inches(3.5))
    shape(s, MSO_SHAPE.OVAL, 10.05, 1.45, 3.7, 3.7, line=ACC, lw=1.5)


def s_glance():
    s = new_slide(MS_SECTION)
    header(
        s,
        "At a glance",
        "One system. Two portals. Zero guesswork.",
        "Everything an operator needs, and a clean self-service view for every resident.",
    )
    stats = [
        ("2", "Portals", ACC),
        ("9", "Modules", ACC2),
        ("147", "Tests passing", GREEN),
        ("100%", "Role-guarded routes", VIOLET),
    ]
    x, w, gap = 0.7, 2.83, 0.35
    for i, (big, label, color) in enumerate(stats):
        cx = x + i * (w + gap)
        card(s, cx, 2.35, w, 1.85)
        shape(s, MSO_SHAPE.RECTANGLE, cx, 2.35, w, 0.08, fill=color)
        txt(
            s,
            cx + 0.32,
            2.7,
            w - 0.6,
            0.9,
            [{"runs": [{"text": big, "size": 38, "bold": True, "color": color, "font": FONT}]}],
        )
        txt(
            s,
            cx + 0.32,
            3.6,
            w - 0.6,
            0.5,
            [{"runs": [{"text": label, "size": 12.5, "color": INK}]}],
        )
    card(s, 0.7, 4.65, 11.93, 1.55)
    txt(
        s,
        1.05,
        4.9,
        11.3,
        1.1,
        bullets(
            [
                "Admins run hostels, rooms, allocations, billing and payments.",
                "Students sign in to view their room, bills, payments and receipts.",
                "One database, one audit trail, one source of truth.",
            ],
            size=13,
            gap=11,
        ),
    )


def s_portals():
    s = new_slide(MS_SECTION)
    header(
        s,
        "Architecture",
        "Two portals, one database",
        "The landing page is decided by the account's role at login.",
    )
    cols = [
        (
            0.7,
            "Admin portal",
            ACC,
            "/admin",
            [
                "Hostels & rooms",
                "Students",
                "Allocations",
                "Fees & structures",
                "Payments & receipts",
                "Reports & audit",
                "Settings",
            ],
        ),
        (
            6.83,
            "Student portal",
            VIOLET,
            "/portal",
            ["Dashboard", "Profile", "My room", "My bills", "My payments", "PDF receipts"],
        ),
    ]
    for x, title, color, tag, items in cols:
        card(s, x, 2.25, 5.8, 4.05)
        shape(s, MSO_SHAPE.RECTANGLE, x, 2.25, 5.8, 0.7, fill=color)
        txt(
            s,
            x + 0.35,
            2.38,
            5.0,
            0.45,
            [{"runs": [{"text": title, "size": 16, "bold": True, "color": WHITE, "font": FONT}]}],
        )
        txt(
            s,
            x + 0.35,
            2.4,
            5.1,
            0.45,
            [
                {
                    "align": PP_ALIGN.RIGHT,
                    "runs": [{"text": tag, "size": 11.5, "bold": True, "color": WHITE}],
                }
            ],
        )
        txt(s, x + 0.35, 3.15, 5.15, 3.0, bullets(items, size=13, gap=11, dot=color))


def s_stack():
    s = new_slide(MS_SECTION)
    header(
        s,
        "Under the hood",
        "Built on a modern, dependable stack",
        "Batteries included, standards-based, easy to run and test.",
    )
    cols, w, h, gx, gy = 3, 3.78, 1.75, 0.3, 0.35
    x0, y0 = 0.7, 2.35
    for i, (icon, name, desc) in enumerate(STACK):
        r, c = divmod(i, cols)
        x = x0 + c * (w + gx)
        y = y0 + r * (h + gy)
        card(s, x, y, w, h)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.32, y + 0.38, 0.78, 0.78, fill=SOFT, radius=0.3)
        txt(
            s,
            x + 0.32,
            y + 0.38,
            0.78,
            0.78,
            [{"align": PP_ALIGN.CENTER, "runs": [{"text": icon, "size": 22, "color": ACC}]}],
            anchor=MSO_ANCHOR.MIDDLE,
        )
        txt(
            s,
            x + 1.3,
            y + 0.4,
            w - 1.55,
            0.4,
            [{"runs": [{"text": name, "size": 14.5, "bold": True, "color": INK, "font": FONT}]}],
        )
        txt(
            s,
            x + 1.3,
            y + 0.86,
            w - 1.55,
            0.6,
            [{"runs": [{"text": desc, "size": 11.5, "color": SUB}]}],
        )


def overview(slide_title, sub, nodes):
    s = new_slide(MS_MAP)
    header(s, "Workflow map", slide_title, sub)
    cols = [0.7, 3.83, 6.96, 10.09]
    w, h = 2.54, 1.0
    r1, r2 = 2.5, 4.75
    pos = {k: (cols[k], r1) for k in range(4)}
    pos.update({4 + j: (cols[3 - j], r2) for j in range(4)})
    for i, name in enumerate(nodes):
        x, y = pos[i]
        card(s, x, y, w, h)
        shape(s, MSO_SHAPE.RECTANGLE, x, y, 0.09, h, fill=ACC)
        txt(
            s,
            x + 0.26,
            y + 0.16,
            w - 0.45,
            h - 0.3,
            [
                {
                    "runs": [{"text": f"STEP {i + 1}", "size": 9, "bold": True, "color": ACC2}],
                    "space_after": 2,
                },
                {
                    "runs": [
                        {
                            "text": name.replace(chr(10), " "),
                            "size": 12.5,
                            "bold": True,
                            "color": INK,
                        }
                    ]
                },
            ],
        )
    for k in range(3):
        shape(
            s, MSO_SHAPE.RIGHT_ARROW, cols[k] + w + 0.05, r1 + h / 2 - 0.14, 0.6, 0.28, fill=SOFT2
        )
    shape(
        s,
        MSO_SHAPE.DOWN_ARROW,
        cols[3] + w / 2 - 0.14,
        r1 + h + 0.08,
        0.28,
        r2 - (r1 + h) - 0.16,
        fill=SOFT2,
    )
    for k in range(3):
        shape(
            s,
            MSO_SHAPE.LEFT_ARROW,
            cols[2 - k] + w + 0.05,
            r2 + h / 2 - 0.14,
            0.6,
            0.28,
            fill=SOFT2,
        )


def step_slide(nodes, i, step):
    s = new_slide(MS_STEP)
    rail(s, nodes, i)
    card(s, 0.7, 1.38, 7.4, 5.62)
    txt(
        s,
        1.05,
        1.7,
        6.7,
        1.1,
        [
            {
                "runs": [
                    {"text": f"STEP {i + 1}   ", "size": 12, "bold": True, "color": ACC2},
                    {"text": step["title"], "size": 21, "bold": True, "color": INK, "font": FONT},
                ],
                "space_after": 7,
            },
            {"runs": [{"text": step["route"], "size": 12, "italic": True, "color": MUTED}]},
        ],
    )
    txt(s, 1.05, 3.0, 6.7, 3.7, bullets(step["bullets"], size=13.5, gap=14, dot=ACC))
    rx, rw, ch, gap = 8.35, 4.28, 1.2, 0.19
    y0 = 1.38 + (5.62 - (4 * ch + 3 * gap)) / 2
    info(s, rx, y0 + 0 * (ch + gap), rw, ch, "Route", step["route"], ACC2)
    info(s, rx, y0 + 1 * (ch + gap), rw, ch, "Data flow", step["data"], ACC)
    info(s, rx, y0 + 2 * (ch + gap), rw, ch, "Code", step["files"], VIOLET)
    info(s, rx, y0 + 3 * (ch + gap), rw, ch, "Guardrails", step["guard"], AMBER)


def s_lifecycle():
    s = new_slide(MS_LIFE)
    header(
        s,
        "End to end",
        "How the two portals connect",
        "The admin creates and bills; the student views and pays down. The cycle repeats monthly.",
    )
    left, right, gap = 0.7, 12.63, 0.22
    n = len(LIFECYCLE)
    w = (right - left - gap * (n - 1)) / n
    y, h = 3.15, 1.85
    for i, (label, role) in enumerate(LIFECYCLE):
        x = left + i * (w + gap)
        color = ROLE_COLOR[role]
        chip(s, x + w / 2 - 0.5, 2.6, 1.0, 0.3, role, color=color)
        card(s, x, y, w, h)
        shape(s, MSO_SHAPE.RECTANGLE, x, y, w, 0.09, fill=color)
        txt(
            s,
            x + 0.15,
            y + 0.32,
            w - 0.3,
            1.35,
            [
                {
                    "align": PP_ALIGN.CENTER,
                    "space_after": 4,
                    "runs": [
                        {"text": f"{i + 1}", "size": 21, "bold": True, "color": color, "font": FONT}
                    ],
                },
                {
                    "align": PP_ALIGN.CENTER,
                    "runs": [
                        {"text": label.replace("\n", " "), "size": 11, "bold": True, "color": INK}
                    ],
                },
            ],
        )
        if i < n - 1:
            shape(
                s,
                MSO_SHAPE.RIGHT_ARROW,
                x + w + 0.01,
                y + h / 2 - 0.12,
                gap - 0.02,
                0.24,
                fill=SOFT2,
            )
    card(s, 0.7, 5.5, 11.93, 1.05)
    txt(
        s,
        1.0,
        5.68,
        11.3,
        0.75,
        [
            {
                "runs": [
                    {"text": "\u21bb  ", "size": 16, "bold": True, "color": ACC},
                    {
                        "text": "Bill \u2192 pay \u2192 reconcile \u2192 report \u2192 "
                        "re-allocate. Deactivating a student frees their bed.",
                        "size": 12.5,
                        "color": INK,
                    },
                ]
            }
        ],
    )


def s_guarantees():
    s = new_slide(MS_GUAR)
    header(
        s,
        "Trust",
        "Why it holds up under real use",
        "Guardrails are enforced in code and in the database \u2014 not just in the UI.",
    )
    cols, w, h, gx, gy = 3, 3.78, 1.7, 0.3, 0.35
    x0, y0 = 0.7, 2.35
    for i, (icon, name, desc) in enumerate(GUARANTEES):
        r, c = divmod(i, cols)
        x = x0 + c * (w + gx)
        y = y0 + r * (h + gy)
        card(s, x, y, w, h)
        shape(s, MSO_SHAPE.RECTANGLE, x, y, w, 0.07, fill=ACC)
        txt(
            s, x + 0.32, y + 0.38, 0.7, 0.55, [{"runs": [{"text": icon, "size": 21, "color": ACC}]}]
        )
        txt(
            s,
            x + 1.0,
            y + 0.38,
            w - 1.25,
            0.42,
            [{"runs": [{"text": name, "size": 14, "bold": True, "color": INK, "font": FONT}]}],
        )
        txt(
            s,
            x + 1.0,
            y + 0.84,
            w - 1.25,
            0.7,
            [{"runs": [{"text": desc, "size": 11, "color": SUB}]}],
        )


def s_thanks():
    s = new_slide(MS_SECTION)
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.18, fill=ACC)
    txt(
        s,
        1.0,
        2.3,
        11.3,
        0.4,
        [{"runs": [{"text": "THANK YOU", "size": 13, "bold": True, "color": ACC}]}],
    )
    txt(
        s,
        1.0,
        2.75,
        11.4,
        1.0,
        [{"runs": [{"text": "Questions?", "size": 40, "bold": True, "color": INK, "font": FONT}]}],
    )
    shape(s, MSO_SHAPE.RECTANGLE, 1.02, 3.85, 2.4, 0.06, fill=ACC2)
    txt(
        s,
        1.02,
        4.05,
        11.0,
        0.5,
        [
            {
                "runs": [
                    {"text": f"{HOSTEL_NAME}  \u00b7  Flask + PostgreSQL", "size": 14, "color": SUB}
                ]
            }
        ],
    )
    card(s, 1.0, 4.8, 11.3, 1.5)
    txt(
        s,
        1.35,
        5.05,
        10.6,
        1.1,
        [
            {
                "runs": [
                    {"text": "python run.py", "size": 14, "bold": True, "color": ACC},
                    {"text": "   \u2192   http://127.0.0.1:5000", "size": 13, "color": SUB},
                ],
                "space_after": 10,
            },
            {
                "runs": [
                    {"text": "python -m pytest", "size": 14, "bold": True, "color": GREEN},
                    {"text": "   \u2192   147 tests", "size": 13, "color": SUB},
                ]
            },
        ],
    )


def build():
    try:
        make_circle_png()
    except Exception as exc:  # image is optional; never block the deck
        print(f"Circle image skipped: {exc}")
    s_welcome()
    s_glance()
    s_portals()
    s_stack()
    overview("Admin portal \u2014 the 8 steps", "From sign-in to reporting.", ADMIN_NODES)
    for i, st in enumerate(ADMIN_STEPS):
        step_slide(ADMIN_NODES, i, st)
    overview("Student portal \u2014 the 8 steps", "From sign-up to receipts.", STUDENT_NODES)
    for i, st in enumerate(STUDENT_STEPS):
        step_slide(STUDENT_NODES, i, st)
    s_lifecycle()
    s_guarantees()
    s_thanks()
    try:
        prs.save(OUT)
        saved = OUT
    except PermissionError:
        saved = OUT.with_name(OUT.stem + "_v2" + OUT.suffix)
        prs.save(saved)
    print(f"Wrote {saved} with {len(prs.slides._sldIdLst)} slides")


if __name__ == "__main__":
    build()
