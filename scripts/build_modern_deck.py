"""Generate a modern, animated workflow deck for the hostel app.

Output: ``Hostel_Modern_Deck.pptx`` in the project root.

Design: dark navy->violet gradient background, translucent "glass" cards,
indigo/violet accents, big type, and a Morph transition on every slide
(PowerPoint's smooth, real-world animation) with a fade fallback so it still
animates in older viewers.  Workflows are also split one-step-per-slide with a
progress rail that lights up node by node, so advancing the deck animates the
flow "pin to pin".
"""

from __future__ import annotations

import contextlib
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUT = PROJECT_ROOT / "Hostel_Modern_Deck.pptx"

# --- palette ---------------------------------------------------------------
BG1, BG2 = "0B1120", "1B1233"  # background gradient
ACC1, ACC2, ACC3 = "4F8DFD", "A855F7", "22D3EE"
GREEN, AMBER, DANGER = "34D399", "FBBF24", "FB7185"
WHITE, SUB, MUTED = "F8FAFC", "93A4C4", "6B7A99"
BORDER = "2B3A5C"
FONT = "Segoe UI"

HOSTEL_NAME = "Sunrise Student Hostel"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

ROLE_COLOR = {"ADMIN": ACC1, "STUDENT": ACC2, "ADMIN + STUDENT": AMBER}


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


# --- transitions -----------------------------------------------------------
def morph(slide):
    sld = slide._element
    for t in sld.findall(qn("p:transition")):
        sld.remove(t)
    xml = (
        "<p:transition "
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" '
        'spd="slow"><mc:AlternateContent>'
        '<mc:Choice Requires="p159"><p159:morph option="byObject"/></mc:Choice>'
        "<mc:Fallback><p:fade/></mc:Fallback>"
        "</mc:AlternateContent></p:transition>"
    )
    ref = sld.find(qn("p:clrMapOvr"))
    if ref is None:
        ref = sld.find(qn("p:cSld"))
    ref.addnext(parse_xml(xml))


# --- primitives ------------------------------------------------------------
def set_alpha(shape, pct):
    spPr = shape._element.spPr
    solid = spPr.find(qn("a:solidFill"))
    if solid is None:
        return
    clr = solid.find(qn("a:srgbClr"))
    if clr is None:
        return
    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(round(pct * 1000))}))


def no_line(shape):
    shape.line.fill.background()


def shape(slide, kind, x, y, w, h, fill=None, line=None, lw=1.0, radius=0.14, alpha=None):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            set_alpha(shp, alpha)
    if line is None:
        no_line(shp)
    else:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        with contextlib.suppress(IndexError, ValueError):
            shp.adjustments[0] = radius
    shp.shadow.inherit = False
    return shp


def grad(slide, kind, x, y, w, h, c1, c2, angle=90.0, radius=0.2):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.gradient()
    with contextlib.suppress(Exception):
        shp.fill.gradient_angle = angle
    stops = shp.fill.gradient_stops
    stops[0].color.rgb = rgb(c1)
    stops[0].position = 0.0
    stops[-1].color.rgb = rgb(c2)
    stops[-1].position = 1.0
    no_line(shp)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        with contextlib.suppress(IndexError, ValueError):
            shp.adjustments[0] = radius
    shp.shadow.inherit = False
    return shp


def _fill(tf, lines):
    tf.clear()
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get("align", PP_ALIGN.LEFT)
        p.level = ln.get("level", 0)
        if "space_after" in ln:
            p.space_after = Pt(ln["space_after"])
        if "space_before" in ln:
            p.space_before = Pt(ln["space_before"])
        if "line_spacing" in ln:
            p.line_spacing = ln["line_spacing"]
        runs = ln.get("runs") or [{"text": ln.get("text", "")}]
        for r in runs:
            run = p.add_run()
            run.text = r.get("text", "")
            f = run.font
            f.size = Pt(r.get("size", ln.get("size", 14)))
            f.bold = r.get("bold", ln.get("bold", False))
            f.italic = r.get("italic", ln.get("italic", False))
            f.name = FONT
            f.color.rgb = rgb(r.get("color", ln.get("color", SUB)))


def txt(slide, x, y, w, h, lines, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    _fill(tf, lines)
    return tb


def stext(shp, lines, anchor=MSO_ANCHOR.MIDDLE, m=0.08):
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(m)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    _fill(tf, lines)


def bullet(items, size=13, color=SUB, gap=8, dot=ACC3):
    return [
        {
            "runs": [
                {"text": "\u25aa  ", "color": dot, "size": size, "bold": True},
                {"text": t, "color": color, "size": size},
            ],
            "space_after": gap,
        }
        for t in items
    ]


# --- slide chrome ----------------------------------------------------------
def new_slide():
    s = prs.slides.add_slide(BLANK)
    morph(s)
    grad(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, BG1, BG2, angle=120)
    shape(s, MSO_SHAPE.OVAL, 10.6, -1.6, 5.2, 5.2, fill=ACC2, alpha=12)
    shape(s, MSO_SHAPE.OVAL, -1.8, 5.0, 5.0, 5.0, fill=ACC1, alpha=12)
    return s


def header(slide, kicker, title, sub=None):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 0.62, 0.12, 0.5, fill=ACC3, radius=0.5)
    txt(
        slide,
        0.7,
        0.56,
        11.5,
        0.3,
        [{"runs": [{"text": kicker.upper(), "size": 11, "bold": True, "color": ACC3}]}],
    )
    txt(
        slide,
        0.7,
        0.86,
        11.9,
        0.7,
        [{"runs": [{"text": title, "size": 30, "bold": True, "color": WHITE}]}],
    )
    if sub:
        txt(slide, 0.72, 1.62, 11.9, 0.45, [{"runs": [{"text": sub, "size": 13.5, "color": SUB}]}])
    grad(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.08, ACC1, ACC2, angle=0)


def card(slide, x, y, w, h, radius=0.06, fill=WHITE, alpha=7):
    shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        x,
        y,
        w,
        h,
        fill=fill,
        alpha=alpha,
        line=BORDER,
        lw=1.0,
        radius=radius,
    )


def chip(slide, x, y, w, h, label, color=ACC1, alpha=16):
    shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        x,
        y,
        w,
        h,
        fill=color,
        alpha=alpha,
        line=color,
        lw=1.0,
        radius=0.5,
    )
    txt(
        slide,
        x,
        y,
        w,
        h,
        [
            {
                "align": PP_ALIGN.CENTER,
                "runs": [{"text": label, "size": 10, "bold": True, "color": WHITE}],
            }
        ],
        anchor=MSO_ANCHOR.MIDDLE,
    )


def rail(slide, nodes, current):
    caption = "STEP-BY-STEP WORKFLOW"
    txt(
        slide,
        0.7,
        0.3,
        8,
        0.26,
        [{"runs": [{"text": caption, "size": 9.5, "bold": True, "color": ACC3}]}],
    )
    n = len(nodes)
    left, right, gap = 0.7, 12.63, 0.13
    w = (right - left - gap * (n - 1)) / n
    y, h = 0.62, 0.46
    for i, name in enumerate(nodes):
        x = left + i * (w + gap)
        if i == current:
            shp = grad(
                slide,
                MSO_SHAPE.ROUNDED_RECTANGLE,
                x,
                y - 0.06,
                w,
                h + 0.12,
                ACC1,
                ACC2,
                angle=0,
                radius=0.35,
            )
            tc = WHITE
        else:
            fill = GREEN if i < current else WHITE
            al = 26 if i < current else 8
            shp = shape(
                slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, alpha=al, radius=0.35
            )
            tc = WHITE if i < current else MUTED
        stext(
            shp,
            [
                {
                    "align": PP_ALIGN.CENTER,
                    "runs": [
                        {"text": f"{i + 1} ", "size": 9.5, "bold": True, "color": tc},
                        {"text": name, "size": 9.5, "bold": (i == current), "color": tc},
                    ],
                }
            ],
        )


def info(slide, x, y, w, h, label, value, color=ACC1):
    card(slide, x, y, w, h)
    grad(slide, MSO_SHAPE.RECTANGLE, x, y + 0.12, 0.08, h - 0.24, color, color, angle=90)
    txt(
        slide,
        x + 0.28,
        y + 0.14,
        w - 0.5,
        h - 0.2,
        [
            {
                "runs": [{"text": label.upper(), "size": 9, "bold": True, "color": color}],
                "space_after": 3,
            },
            {"runs": [{"text": value, "size": 12, "color": WHITE}]},
        ],
    )


# --- content ---------------------------------------------------------------
ADMIN_NODES = [
    "Sign in",
    "Dashboard",
    "Hostels & Rooms",
    "Students",
    "Allocations",
    "Fees",
    "Payments",
    "Reports & Audit",
]
STUDENT_NODES = [
    "Sign up",
    "Sign in",
    "Dashboard",
    "Profile",
    "Allocations",
    "Fees",
    "Payments",
    "Receipts",
]

ADMIN_STEPS = [
    {
        "title": "Admin signs in",
        "route": "/auth/login \u2192 /admin/",
        "bullets": [
            "Email + password verified against the stored hash.",
            "A Flask-Login session opens and last_login_at is stamped.",
            "Role decides the landing page: admin \u2192 dashboard.",
        ],
        "files": "app/auth/routes.py",
        "data": "credentials \u2192 session",
        "guard": "rate-limited \u00b7 no enumeration \u00b7 role redirect",
    },
    {
        "title": "Operational dashboard",
        "route": "/admin/  (admin.dashboard)",
        "bullets": [
            "Live KPIs: students, rooms, capacity, occupancy.",
            "Monthly collection, outstanding and overdue count.",
            "Charts: by room type, by hostel, 6-month trend.",
        ],
        "files": "admin/routes.py \u00b7 services/dashboard_service.py",
        "data": "SQL aggregates \u2192 cards + charts",
        "guard": "admin_required",
    },
    {
        "title": "Set up hostels & rooms",
        "route": "/hostels \u00b7 /rooms",
        "bullets": [
            "Create a hostel, then add rooms (type, capacity, base fee).",
            "Capacity can never drop below current occupancy.",
            "Rooms with history can't be deleted \u2014 make unavailable.",
        ],
        "files": "rooms/routes.py \u00b7 rooms/forms.py",
        "data": "form \u2192 Hostel / Room rows",
        "guard": "capacity \u2265 occupancy \u00b7 unique number",
    },
    {
        "title": "Onboard students",
        "route": "/students",
        "bullets": [
            "Create a student + optional portal login in one step.",
            "Filter by course, year, status, allocation, fee state.",
            "Deactivating automatically releases the active room.",
        ],
        "files": "students/routes.py",
        "data": "profile \u2192 Student (+ User)",
        "guard": "unique email & roll no \u00b7 deactivation frees bed",
    },
    {
        "title": "Allocate beds",
        "route": "/allocations",
        "bullets": [
            "Auto, manual, transfer and release operations.",
            "Auto picks the lowest free bed in the first eligible room.",
            "Concurrency-safe: SKIP LOCKED + unique indexes + trigger.",
        ],
        "files": "allocations/routes.py \u00b7 services/allocation_service.py",
        "data": "student + room \u2192 active allocation",
        "guard": "no double-booking \u00b7 retries on contention",
    },
    {
        "title": "Generate monthly bills",
        "route": "/fees  (+ /fees/structures)",
        "bullets": [
            "Structures per room type set the default amounts.",
            "Generate for a period (YYYY-MM); duplicates skipped.",
            "Late fee after grace; status paid / partial / overdue / pending.",
        ],
        "files": "fees/routes.py \u00b7 services/fee_service.py",
        "data": "rooms + structures + period \u2192 StudentFee rows",
        "guard": "Decimal money \u00b7 idempotent generation",
    },
    {
        "title": "Record payments & receipts",
        "route": "/payments",
        "bullets": [
            "Idempotency key prevents duplicate payments.",
            "Fee row locked (FOR UPDATE); overpayment rejected.",
            "Immutable \u2014 corrections are linked reversals; PDF receipt.",
        ],
        "files": "payments/routes.py \u00b7 services/payment_service.py",
        "data": "payment \u2192 Payment row \u2192 balance update",
        "guard": "idempotent \u00b7 immutable \u00b7 receipt PDF",
    },
    {
        "title": "Reports, audit & settings",
        "route": "/reports \u00b7 /admin/audit \u00b7 /admin/settings",
        "bullets": [
            "Slug-based reports with CSV export + occupancy/financial PDFs.",
            "Audit log records who changed what.",
            "Settings: hostel identity, currency, late-fee policy.",
        ],
        "files": "reports/routes.py \u00b7 admin/routes.py",
        "data": "DB \u2192 CSV / PDF \u00b7 every write \u2192 audit",
        "guard": "admin_required \u00b7 immutable audit trail",
    },
]

STUDENT_STEPS = [
    {
        "title": "Create a portal account",
        "route": "/auth/signup",
        "bullets": [
            "Signup creates the User (role = student) + Student profile.",
            "Email and roll number must be unique.",
            "New accounts link straight to the student portal.",
        ],
        "files": "app/auth/routes.py",
        "data": "signup \u2192 User + Student",
        "guard": "rate-limited \u00b7 unique email & roll no",
    },
    {
        "title": "Student signs in",
        "route": "/auth/login \u2192 /portal",
        "bullets": [
            "Same login endpoint; the role picks the destination.",
            "Students always land on the portal dashboard.",
        ],
        "files": "app/auth/routes.py",
        "data": "credentials \u2192 session",
        "guard": "rate-limited \u00b7 role redirect",
    },
    {
        "title": "Personal dashboard",
        "route": "/portal",
        "bullets": [
            "Current room allocation and recent payments.",
            "Upcoming due bill and totals: billed, paid, outstanding.",
            "Everything scoped to the logged-in student.",
        ],
        "files": "students/portal.py",
        "data": "own records \u2192 summary",
        "guard": "student_required \u00b7 self-scoped",
    },
    {
        "title": "Profile self-service",
        "route": "/portal/profile",
        "bullets": [
            "View the profile; edit phone, address, emergency contact.",
            "Identity fields stay admin-controlled.",
        ],
        "files": "students/portal.py",
        "data": "limited form \u2192 Student",
        "guard": "editable fields restricted",
    },
    {
        "title": "My room",
        "route": "/portal/allocations",
        "bullets": [
            "Current bed plus full allocation history.",
            "Read-only \u2014 allocation is an admin action.",
        ],
        "files": "students/portal.py",
        "data": "own allocations \u2192 list",
        "guard": "self-scoped",
    },
    {
        "title": "My bills",
        "route": "/portal/fees",
        "bullets": [
            "Monthly bills with amount, paid, balance, status.",
            "Per-bill detail shows the breakdown and due date.",
            "Opening someone else's bill returns 404.",
        ],
        "files": "students/portal.py",
        "data": "own fees \u2192 list / detail",
        "guard": "IDOR-safe (404 on others)",
    },
    {
        "title": "My payments",
        "route": "/portal/payments",
        "bullets": [
            "History of payments recorded against the account.",
            "Read-only \u2014 recording is an admin action.",
        ],
        "files": "students/portal.py",
        "data": "own payments \u2192 list",
        "guard": "self-scoped",
    },
    {
        "title": "Download receipts",
        "route": "/portal/receipts/<id>",
        "bullets": [
            "Download a PDF receipt for any of your own payments.",
            "A receipt that isn't yours returns 404.",
        ],
        "files": "students/portal.py",
        "data": "payment id \u2192 own PDF",
        "guard": "IDOR-safe (404 on others)",
    },
]

STACK = [
    ("\u2699", "Flask 3", "App factory + blueprints"),
    ("\U0001f5c4", "PostgreSQL", "SQLAlchemy + Alembic"),
    ("\U0001f510", "Flask-Login", "Role-based sessions"),
    ("\U0001f310", "Bootstrap 5", "Responsive UI"),
    ("\U0001f4c4", "ReportLab", "PDF receipts & reports"),
    ("\u2713", "pytest", "147 tests passing"),
]

GUARANTEES = [
    ("\U0001f510", "Role-checked routes", "admin_required / student_required on every view"),
    ("\U0001f5c3", "DB-level integrity", "partial unique indexes + capacity triggers"),
    ("\U0001f4b3", "Immutable payments", "corrections are linked reversal records"),
    ("\U0001f464", "Self-scoped data", "students never see other records (404)"),
    ("\U0001f522", "Decimal money", "no floating-point drift in billing"),
    ("\U0001f4cb", "Full audit trail", "every state change is logged"),
]

LIFECYCLE = [
    ("Set up hostels\n& rooms", "ADMIN"),
    ("Enroll\nstudents", "ADMIN + STUDENT"),
    ("Allocate\nbeds", "ADMIN"),
    ("Generate\nmonthly fees", "ADMIN"),
    ("Record\npayments", "ADMIN"),
    ("View bills &\nreceipts", "STUDENT"),
    ("Reports, audit\n& deactivate", "ADMIN"),
]

FLOW = ["Setup", "Rooms", "Students", "Allocate", "Bill", "Pay", "Report"]


# --- slides ----------------------------------------------------------------
def s_title():
    s = new_slide()
    txt(
        s,
        0.9,
        1.7,
        11.5,
        0.4,
        [
            {
                "runs": [
                    {
                        "text": "FLASK  \u00b7  POSTGRESQL  \u00b7  ROLE-BASED PORTALS",
                        "size": 12,
                        "bold": True,
                        "color": ACC3,
                    }
                ]
            }
        ],
    )
    txt(
        s,
        0.9,
        2.15,
        11.8,
        1.6,
        [
            {
                "runs": [{"text": "Hostel Management", "size": 52, "bold": True, "color": WHITE}],
                "space_after": 0,
                "line_spacing": 0.95,
            },
            {
                "runs": [{"text": "System", "size": 52, "bold": True, "color": ACC2}],
                "line_spacing": 0.95,
            },
        ],
    )
    txt(
        s,
        0.92,
        4.05,
        10.5,
        0.6,
        [
            {
                "runs": [
                    {
                        "text": "The complete admin + student portal \u2014 set up, allocate, "
                        "bill, collect and report.",
                        "size": 15,
                        "color": SUB,
                    }
                ]
            }
        ],
    )
    x = 0.9
    for label in FLOW:
        w = 0.32 + 0.13 * len(label)
        grad(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 5.0, w, 0.42, ACC1, ACC2, angle=0, radius=0.5)
        txt(
            s,
            x,
            5.0,
            w,
            0.42,
            [
                {
                    "align": PP_ALIGN.CENTER,
                    "runs": [{"text": label, "size": 10.5, "bold": True, "color": WHITE}],
                }
            ],
            anchor=MSO_ANCHOR.MIDDLE,
        )
        x += w + 0.12
    txt(
        s,
        0.9,
        6.3,
        11,
        0.4,
        [
            {
                "runs": [
                    {
                        "text": f"{HOSTEL_NAME}   |   a pin-to-pin workflow walkthrough",
                        "size": 12,
                        "color": MUTED,
                    }
                ]
            }
        ],
    )


def s_glance():
    s = new_slide()
    header(
        s,
        "At a glance",
        "One system. Two portals. Zero guesswork.",
        "Everything an operator needs, and a clean self-service view for every resident.",
    )
    stats = [
        ("2", "Portals", ACC1),
        ("9", "Modules", ACC3),
        ("147", "Tests passing", GREEN),
        ("100%", "Role-guarded routes", ACC2),
    ]
    x, w, gap = 0.7, 2.83, 0.35
    for i, (big, label, color) in enumerate(stats):
        cx = x + i * (w + gap)
        card(s, cx, 2.5, w, 1.9)
        grad(s, MSO_SHAPE.RECTANGLE, cx, 2.5, w, 0.09, color, color, angle=0)
        txt(
            s,
            cx + 0.3,
            2.8,
            w - 0.6,
            1.0,
            [{"runs": [{"text": big, "size": 40, "bold": True, "color": color}]}],
        )
        txt(
            s,
            cx + 0.3,
            3.75,
            w - 0.6,
            0.5,
            [{"runs": [{"text": label, "size": 13, "color": WHITE}]}],
        )
    card(s, 0.7, 4.9, 11.93, 1.4)
    txt(
        s,
        1.05,
        5.15,
        11.3,
        1.0,
        bullet(
            [
                "Admins run hostels, rooms, allocations, billing and payments.",
                "Students sign in to view their room, bills, payments and receipts.",
                "One database, one audit trail, one source of truth.",
            ],
            size=13.5,
            gap=10,
        ),
    )


def s_portals():
    s = new_slide()
    header(
        s,
        "Architecture",
        "Two portals, one database",
        "The landing page is decided by the account's role at login.",
    )
    cols = [
        (
            0.7,
            "Admin portal",
            ACC1,
            "/admin",
            [
                "Hostels & rooms",
                "Students",
                "Allocations",
                "Fees & structures",
                "Payments & receipts",
                "Reports & audit",
                "Settings",
            ],
        ),
        (
            6.83,
            "Student portal",
            ACC2,
            "/portal",
            ["Dashboard", "Profile", "My room", "My bills", "My payments", "PDF receipts"],
        ),
    ]
    for x, title, color, tag, items in cols:
        card(s, x, 2.25, 5.8, 4.05)
        grad(s, MSO_SHAPE.RECTANGLE, x, 2.25, 5.8, 0.7, color, color, angle=0)
        txt(
            s,
            x + 0.35,
            2.38,
            5.0,
            0.45,
            [{"runs": [{"text": title, "size": 17, "bold": True, "color": WHITE}]}],
        )
        txt(
            s,
            x + 0.35,
            2.4,
            5.1,
            0.45,
            [
                {
                    "align": PP_ALIGN.RIGHT,
                    "runs": [{"text": tag, "size": 11.5, "bold": True, "color": WHITE}],
                }
            ],
        )
        txt(s, x + 0.35, 3.15, 5.15, 3.0, bullet(items, size=13.5, gap=11))


def s_stack():
    s = new_slide()
    header(
        s,
        "Under the hood",
        "Built on a modern, boring-in-a-good-way stack",
        "Batteries included, standards-based, easy to run and test.",
    )
    cols, w, h, gx, gy = 3, 3.78, 1.75, 0.3, 0.35
    x0, y0 = 0.7, 2.35
    for i, (icon, name, desc) in enumerate(STACK):
        r, c = divmod(i, cols)
        x = x0 + c * (w + gx)
        y = y0 + r * (h + gy)
        card(s, x, y, w, h)
        shape(s, MSO_SHAPE.OVAL, x + 0.35, y + 0.4, 0.8, 0.8, fill=ACC1, alpha=18)
        txt(
            s,
            x + 0.35,
            y + 0.4,
            0.8,
            0.8,
            [{"align": PP_ALIGN.CENTER, "runs": [{"text": icon, "size": 22, "color": WHITE}]}],
            anchor=MSO_ANCHOR.MIDDLE,
        )
        txt(
            s,
            x + 1.35,
            y + 0.42,
            w - 1.6,
            0.4,
            [{"runs": [{"text": name, "size": 15, "bold": True, "color": WHITE}]}],
        )
        txt(
            s,
            x + 1.35,
            y + 0.88,
            w - 1.6,
            0.6,
            [{"runs": [{"text": desc, "size": 12, "color": SUB}]}],
        )


def overview(slide_title, sub, nodes):
    s = new_slide()
    header(s, "Workflow map", slide_title, sub)
    cols = [0.7, 3.83, 6.96, 10.09]
    w, h = 2.54, 1.05
    r1, r2 = 2.5, 4.75
    pos = {k: (cols[k], r1) for k in range(4)}
    pos.update({4 + j: (cols[3 - j], r2) for j in range(4)})
    for i, name in enumerate(nodes):
        x, y = pos[i]
        card(s, x, y, w, h)
        grad(s, MSO_SHAPE.RECTANGLE, x, y, 0.1, h, ACC1, ACC2, angle=90)
        txt(
            s,
            x + 0.28,
            y + 0.16,
            w - 0.5,
            h - 0.3,
            [
                {
                    "runs": [{"text": f"STEP {i + 1}", "size": 9, "bold": True, "color": ACC3}],
                    "space_after": 2,
                },
                {
                    "runs": [
                        {
                            "text": name.replace(chr(10), " "),
                            "size": 13,
                            "bold": True,
                            "color": WHITE,
                        }
                    ]
                },
            ],
        )
    for k in range(3):
        shape(
            s,
            MSO_SHAPE.RIGHT_ARROW,
            cols[k] + w + 0.05,
            r1 + h / 2 - 0.15,
            0.6,
            0.3,
            fill=ACC1,
            alpha=40,
        )
    shape(
        s,
        MSO_SHAPE.DOWN_ARROW,
        cols[3] + w / 2 - 0.15,
        r1 + h + 0.08,
        0.3,
        r2 - (r1 + h) - 0.16,
        fill=ACC2,
        alpha=40,
    )
    for k in range(3):
        shape(
            s,
            MSO_SHAPE.LEFT_ARROW,
            cols[2 - k] + w + 0.05,
            r2 + h / 2 - 0.15,
            0.6,
            0.3,
            fill=ACC1,
            alpha=40,
        )


def step_slide(nodes, i, step):
    s = new_slide()
    rail(s, nodes, i)
    card(s, 0.7, 1.4, 7.4, 5.6)
    txt(
        s,
        1.05,
        1.72,
        6.7,
        1.1,
        [
            {
                "runs": [
                    {"text": f"STEP {i + 1}   ", "size": 12, "bold": True, "color": ACC3},
                    {"text": step["title"], "size": 22, "bold": True, "color": WHITE},
                ],
                "space_after": 7,
            },
            {"runs": [{"text": step["route"], "size": 12, "italic": True, "color": MUTED}]},
        ],
    )
    txt(s, 1.05, 3.05, 6.7, 3.7, bullet(step["bullets"], size=14, gap=14, dot=ACC2))
    rx, rw, ch, gap = 8.35, 4.28, 1.2, 0.19
    y0 = 1.4 + (5.6 - (4 * ch + 3 * gap)) / 2
    info(s, rx, y0 + 0 * (ch + gap), rw, ch, "Route", step["route"], ACC3)
    info(s, rx, y0 + 1 * (ch + gap), rw, ch, "Data flow", step["data"], ACC1)
    info(s, rx, y0 + 2 * (ch + gap), rw, ch, "Code", step["files"], ACC2)
    info(s, rx, y0 + 3 * (ch + gap), rw, ch, "Guardrails", step["guard"], AMBER)


def s_lifecycle():
    s = new_slide()
    header(
        s,
        "End to end",
        "How the two portals connect",
        "The admin creates and bills; the student views and pays down. The cycle repeats monthly.",
    )
    left, right, gap = 0.7, 12.63, 0.22
    n = len(LIFECYCLE)
    w = (right - left - gap * (n - 1)) / n
    y, h = 3.15, 1.9
    for i, (label, role) in enumerate(LIFECYCLE):
        x = left + i * (w + gap)
        color = ROLE_COLOR[role]
        chip(s, x + w / 2 - 0.5, 2.6, 1.0, 0.3, role, color=color, alpha=25)
        card(s, x, y, w, h)
        grad(s, MSO_SHAPE.RECTANGLE, x, y, w, 0.1, color, color, angle=0)
        txt(
            s,
            x + 0.15,
            y + 0.35,
            w - 0.3,
            1.4,
            [
                {
                    "align": PP_ALIGN.CENTER,
                    "space_after": 4,
                    "runs": [{"text": f"{i + 1}", "size": 22, "bold": True, "color": color}],
                },
                {
                    "align": PP_ALIGN.CENTER,
                    "runs": [
                        {
                            "text": label.replace("\n", " "),
                            "size": 11.5,
                            "bold": True,
                            "color": WHITE,
                        }
                    ],
                },
            ],
        )
        if i < n - 1:
            shape(
                s,
                MSO_SHAPE.RIGHT_ARROW,
                x + w + 0.01,
                y + h / 2 - 0.13,
                gap - 0.02,
                0.26,
                fill=ACC1,
                alpha=45,
            )
    card(s, 0.7, 5.5, 11.93, 1.05)
    txt(
        s,
        1.0,
        5.68,
        11.3,
        0.75,
        [
            {
                "runs": [
                    {"text": "\u21bb  ", "size": 16, "bold": True, "color": ACC3},
                    {
                        "text": "Bill \u2192 pay \u2192 reconcile \u2192 report \u2192 "
                        "re-allocate. Deactivating a student frees their bed.",
                        "size": 13,
                        "color": WHITE,
                    },
                ]
            }
        ],
    )


def s_guarantees():
    s = new_slide()
    header(
        s,
        "Trust",
        "Why it holds up under real use",
        "Guardrails are enforced in code and in the database \u2014 not just in the UI.",
    )
    cols, w, h, gx, gy = 3, 3.78, 1.75, 0.3, 0.35
    x0, y0 = 0.7, 2.35
    for i, (icon, name, desc) in enumerate(GUARANTEES):
        r, c = divmod(i, cols)
        x = x0 + c * (w + gx)
        y = y0 + r * (h + gy)
        card(s, x, y, w, h)
        grad(s, MSO_SHAPE.RECTANGLE, x, y, w, 0.08, ACC1, ACC2, angle=0)
        txt(
            s,
            x + 0.35,
            y + 0.42,
            0.7,
            0.6,
            [{"runs": [{"text": icon, "size": 22, "color": WHITE}]}],
        )
        txt(
            s,
            x + 1.05,
            y + 0.42,
            w - 1.3,
            0.45,
            [{"runs": [{"text": name, "size": 14.5, "bold": True, "color": WHITE}]}],
        )
        txt(
            s,
            x + 1.05,
            y + 0.9,
            w - 1.3,
            0.7,
            [{"runs": [{"text": desc, "size": 11.5, "color": SUB}]}],
        )


def s_closing():
    s = new_slide()
    txt(
        s,
        0.9,
        1.9,
        11.5,
        0.5,
        [{"runs": [{"text": "READY TO RUN", "size": 12, "bold": True, "color": ACC3}]}],
    )
    txt(
        s,
        0.9,
        2.35,
        11.8,
        1.2,
        [
            {
                "runs": [
                    {
                        "text": "Start the app. Present the flows. Ship it.",
                        "size": 34,
                        "bold": True,
                        "color": WHITE,
                    }
                ]
            }
        ],
    )
    card(s, 0.9, 3.9, 11.5, 2.05)
    txt(
        s,
        1.25,
        4.15,
        10.9,
        1.6,
        [
            {
                "runs": [
                    {"text": "python run.py", "size": 15, "bold": True, "color": ACC3},
                    {"text": "   \u2192   http://127.0.0.1:5000", "size": 14, "color": SUB},
                ],
                "space_after": 12,
            },
            {
                "runs": [
                    {"text": "python -m pytest", "size": 15, "bold": True, "color": GREEN},
                    {"text": "   \u2192   147 tests", "size": 14, "color": SUB},
                ],
                "space_after": 12,
            },
            {
                "runs": [
                    {"text": "python -m ruff check .", "size": 15, "bold": True, "color": ACC2},
                    {"text": "   \u2192   clean lint", "size": 14, "color": SUB},
                ]
            },
        ],
    )
    txt(
        s,
        0.9,
        6.35,
        11,
        0.4,
        [
            {
                "runs": [
                    {
                        "text": "Press \u2192 to advance each workflow step \u2014 Morph animates "
                        "the transition between slides.",
                        "size": 12,
                        "color": MUTED,
                    }
                ]
            }
        ],
    )


def build():
    s_title()
    s_glance()
    s_portals()
    s_stack()
    overview("Admin portal \u2014 the 8 steps", "From sign-in to reporting.", ADMIN_NODES)
    for i, st in enumerate(ADMIN_STEPS):
        step_slide(ADMIN_NODES, i, st)
    overview("Student portal \u2014 the 8 steps", "From sign-up to receipts.", STUDENT_NODES)
    for i, st in enumerate(STUDENT_STEPS):
        step_slide(STUDENT_NODES, i, st)
    s_lifecycle()
    s_guarantees()
    s_closing()
    prs.save(OUT)
    print(f"Wrote {OUT} with {len(prs.slides._sldIdLst)} slides")


if __name__ == "__main__":
    build()
